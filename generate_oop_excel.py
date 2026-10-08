"""Generate OOP Summary Excel file for the Lost & Found System project."""

import openpyxl
from openpyxl.styles import (
    Font, PatternFill, Alignment, Border, Side
)
from openpyxl.utils import get_column_letter

wb = openpyxl.Workbook()

# ─── Color palette ───────────────────────────────────────────────────────────
NAVY       = "18324B"
TEAL       = "087F78"
TEAL_SOFT  = "E4F4F2"
AMBER      = "FFF3D9"
WHITE      = "FFFFFF"
LIGHT_GRAY = "F4F7FA"
LINE       = "DFE7ED"

def hdr_font(color="FFFFFF", size=10, bold=True):
    return Font(name="Calibri", bold=bold, color=color, size=size)

def cell_font(bold=False, size=10, color="000000"):
    return Font(name="Calibri", bold=bold, size=size, color=color)

def fill(hex_color):
    return PatternFill("solid", fgColor=hex_color)

def border():
    thin = Side(style="thin", color=LINE)
    return Border(left=thin, right=thin, top=thin, bottom=thin)

def wrap_center(horizontal="center", vertical="center", wrap=False):
    return Alignment(horizontal=horizontal, vertical=vertical, wrap_text=wrap)

def style_header_row(ws, row, bg, fg="FFFFFF", height=22):
    ws.row_dimensions[row].height = height
    for cell in ws[row]:
        cell.font      = hdr_font(color=fg)
        cell.fill      = fill(bg)
        cell.alignment = wrap_center("center", "center")
        cell.border    = border()

def style_data_row(ws, row_idx, even, height=18):
    ws.row_dimensions[row_idx].height = height
    bg = LIGHT_GRAY if even else WHITE
    for cell in ws.iter_rows(min_row=row_idx, max_row=row_idx):
        for c in cell:
            c.fill      = fill(bg)
            c.font      = cell_font()
            c.alignment = wrap_center("left", "center", wrap=True)
            c.border    = border()

def set_col_widths(ws, widths):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

def title_row(ws, text, merge_to, row=1, bg=NAVY):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=merge_to)
    cell = ws.cell(row=row, column=1, value=text)
    cell.font      = Font(name="Calibri", bold=True, size=13, color="FFFFFF")
    cell.fill      = fill(bg)
    cell.alignment = wrap_center("center", "center")
    ws.row_dimensions[row].height = 28


# ══════════════════════════════════════════════════════════════════════════════
# SHEET 1 — CLASSES
# ══════════════════════════════════════════════════════════════════════════════
ws_classes = wb.active
ws_classes.title = "Classes"

title_row(ws_classes, "Lost & Found System — OOP Classes", 5)

headers = ["#", "Class Name", "File", "Parent Class", "Description"]
for col, h in enumerate(headers, 1):
    ws_classes.cell(row=2, column=col, value=h)
style_header_row(ws_classes, 2, TEAL)

classes_data = [
    (1,  "User",                "models/user.py",                    "object (dataclass)",  "Frozen dataclass — immutable identity object with role and display_name"),
    (2,  "PublicUser",          "models/user.py",                    "User",                "Unauthenticated visitor; no login flow"),
    (3,  "Admin",               "models/admin.py",                   "User",                "The only authenticated role; handles PBKDF2 password hashing and verification"),
    (4,  "SchoolPerson",        "models/person.py",                  "ABC (Abstract)",      "Abstract base class for Student, Teacher, and UtilityStaff"),
    (5,  "Student",             "models/person.py",                  "SchoolPerson",        "Concrete school person — stores grade/course and section"),
    (6,  "Teacher",             "models/person.py",                  "SchoolPerson",        "Concrete school person — stores department"),
    (7,  "UtilityStaff",        "models/person.py",                  "SchoolPerson",        "Concrete school person — stores assigned area"),
    (8,  "Item",                "models/item.py",                    "object (dataclass)",  "Found-property entity; exposes public_record() vs to_record() boundary"),
    (9,  "Finder",              "models/finder.py",                  "object (dataclass)",  "Finder interview record; composed with a SchoolPerson"),
    (10, "Claim",               "models/claim.py",                   "object (dataclass)",  "Claim request entity; composed with a SchoolPerson as claimant"),
    (11, "DuplicateClaimError", "database/database.py",              "ValueError",          "Raised when the same school ID has an active claim for the same item"),
    (12, "ItemUnavailableError","database/database.py",              "ValueError",          "Raised on invalid item status transitions"),
    (13, "DatabaseManager",     "database/database.py",              "object",              "Owns schema creation, demo seeding, and all SQLite operations"),
    (14, "AdminController",     "controllers/admin_controller.py",   "object",              "Handles admin authentication and session logging"),
    (15, "ItemController",      "controllers/item_controller.py",    "object",              "Handles found-item catalog and admin maintenance"),
    (16, "ClaimController",     "controllers/claim_controller.py",   "object",              "Handles claim submission and admin review"),
    (17, "LostFoundSystem",     "controllers/system_controller.py",  "object",              "Facade — wires DatabaseManager and all three controllers"),
    (18, "LostFoundApp",        "views/app.py",                      "tk.Tk",               "Main Tkinter window; manages view navigation"),
    (19, "AdminLoginView",      "views/admin_login.py",              "ttk.Frame",           "Admin sign-in screen with username/password form"),
    (20, "PublicHomeView",      "views/public_views.py",             "ttk.Frame",           "Public item catalog browsing screen"),
    (21, "AdminDashboardView",  "views/admin_views.py",              "ttk.Frame",           "Admin management panel"),
    (22, "ScrollableFrame",     "views/ui.py",                       "ttk.Frame",           "Reusable scrollable canvas container with mousewheel support"),
]

for i, row in enumerate(classes_data, 3):
    for col, val in enumerate(row, 1):
        ws_classes.cell(row=i, column=col, value=val)
    style_data_row(ws_classes, i, i % 2 == 0)

set_col_widths(ws_classes, [5, 22, 36, 22, 52])


# ══════════════════════════════════════════════════════════════════════════════
# SHEET 2 — OBJECTS
# ══════════════════════════════════════════════════════════════════════════════
ws_objects = wb.create_sheet("Objects")

title_row(ws_objects, "Lost & Found System — Runtime Objects (Instances)", 5, bg=NAVY)

headers = ["#", "Object / Variable", "Type (Class)", "Created In", "Description"]
for col, h in enumerate(headers, 1):
    ws_objects.cell(row=2, column=col, value=h)
style_header_row(ws_objects, 2, TEAL)

objects_data = [
    (1,  "self.system",        "LostFoundSystem",   "LostFoundApp.__init__()",             "Main system facade wiring all controllers"),
    (2,  "self.database",      "DatabaseManager",   "LostFoundSystem.__init__()",          "SQLite persistence layer instance"),
    (3,  "self.items",         "ItemController",    "LostFoundSystem.__init__()",          "Controller for item-related operations"),
    (4,  "self.claims",        "ClaimController",   "LostFoundSystem.__init__()",          "Controller for claim-related operations"),
    (5,  "self.admin",         "AdminController",   "LostFoundSystem.__init__()",          "Controller for admin auth and logging"),
    (6,  "self.current_admin", "Admin | None",      "LostFoundApp",                        "Tracks the currently logged-in admin; None when public"),
    (7,  "self._active_view",  "ttk.Frame | None",  "LostFoundApp",                        "Reference to the currently displayed screen"),
    (8,  "admin",              "Admin",             "DatabaseManager.authenticate_admin()","Returned on successful login"),
    (9,  "person",             "Student / Teacher / UtilityStaff", "make_school_person()", "Created by factory based on role string"),
    (10, "self.username_var",  "tk.StringVar",      "AdminLoginView.__init__()",           "Bound to the username Entry widget"),
    (11, "self.password_var",  "tk.StringVar",      "AdminLoginView.__init__()",           "Bound to the password Entry widget"),
    (12, "self.container",     "ttk.Frame",         "LostFoundApp.__init__()",             "Root frame that holds all swappable views"),
]

for i, row in enumerate(objects_data, 3):
    for col, val in enumerate(row, 1):
        ws_objects.cell(row=i, column=col, value=val)
    style_data_row(ws_objects, i, i % 2 == 0)

set_col_widths(ws_objects, [5, 22, 28, 36, 46])


# ══════════════════════════════════════════════════════════════════════════════
# SHEET 3 — METHODS
# ══════════════════════════════════════════════════════════════════════════════
ws_methods = wb.create_sheet("Methods")

title_row(ws_methods, "Lost & Found System — Methods", 5, bg=NAVY)

headers = ["#", "Class", "Method", "Type", "Description"]
for col, h in enumerate(headers, 1):
    ws_methods.cell(row=2, column=col, value=h)
style_header_row(ws_methods, 2, TEAL)

methods_data = [
    # PublicUser
    (1,  "PublicUser",          "__init__()",                                    "Instance",       "Sets role='Public User', display_name='Visitor'"),
    # Admin
    (2,  "Admin",               "__init__(username, admin_id)",                  "Instance",       "Calls super().__init__() with role='Admin'"),
    (3,  "Admin",               "hash_password(password, salt)",                 "Class method",   "Returns (salt, PBKDF2_hash) tuple"),
    (4,  "Admin",               "verify_password(password, salt, expected_hash)","Class method",   "Constant-time password comparison via hmac"),
    # SchoolPerson
    (5,  "SchoolPerson",        "__init__(full_name, school_id, contact_number)","Instance",       "Stores protected _full_name, _school_id, _contact_number"),
    (6,  "SchoolPerson",        "full_name",                                     "@property",      "Returns _full_name"),
    (7,  "SchoolPerson",        "school_id",                                     "@property",      "Returns _school_id"),
    (8,  "SchoolPerson",        "contact_number",                                "@property",      "Returns _contact_number"),
    (9,  "SchoolPerson",        "role",                                          "@abstractproperty","Must return role string in every subclass"),
    (10, "SchoolPerson",        "affiliation",                                   "@abstractproperty","Must return affiliation string in every subclass"),
    (11, "SchoolPerson",        "section",                                       "@property",      "Returns None by default; overridden in Student"),
    (12, "SchoolPerson",        "to_record()",                                   "Instance",       "Returns dict of all person fields (polymorphic)"),
    # Student
    (13, "Student",             "__init__(..., grade_course, section)",          "Instance",       "Calls super().__init__(), stores grade/course and section"),
    (14, "Student",             "role",                                          "@property",      "Returns 'Student'"),
    (15, "Student",             "affiliation",                                   "@property",      "Returns _grade_course"),
    (16, "Student",             "section",                                       "@property",      "Returns _section (overrides parent)"),
    # Teacher
    (17, "Teacher",             "__init__(..., department)",                     "Instance",       "Calls super().__init__(), stores department"),
    (18, "Teacher",             "role",                                          "@property",      "Returns 'Teacher'"),
    (19, "Teacher",             "affiliation",                                   "@property",      "Returns _department"),
    # UtilityStaff
    (20, "UtilityStaff",        "__init__(..., assigned_area)",                  "Instance",       "Calls super().__init__(), stores assigned area"),
    (21, "UtilityStaff",        "role",                                          "@property",      "Returns 'School Utility Personnel'"),
    (22, "UtilityStaff",        "affiliation",                                   "@property",      "Returns _assigned_area"),
    # Item
    (23, "Item",                "to_record()",                                   "Instance",       "Returns full dict for DB insert/update"),
    (24, "Item",                "public_record()",                               "Instance",       "Returns safe public subset — hides finder/storage/notes"),
    (25, "Item",                "from_mapping(data)",                            "Class method",   "Builds Item from a DB row mapping"),
    # Finder
    (26, "Finder",              "to_record()",                                   "Instance",       "Merges person.to_record() + interview answer fields"),
    # Claim
    (27, "Claim",               "to_record()",                                   "Instance",       "Merges claimant.to_record() + claim fields"),
    (28, "Claim",               "from_mapping(data)",                            "Class method",   "Returns plain dict for display"),
    # DatabaseManager
    (29, "DatabaseManager",     "__init__(db_path, seed_demo)",                  "Instance",       "Sets up path, schema, admin seed, demo items"),
    (30, "DatabaseManager",     "_connection()",                                 "Context manager","Yields sqlite3.Connection with auto commit/rollback"),
    (31, "DatabaseManager",     "_now()",                                        "Static method",  "Returns current ISO datetime string"),
    (32, "DatabaseManager",     "_row_dict(row)",                                "Static method",  "Converts sqlite3.Row → dict"),
    (33, "DatabaseManager",     "_initialize_schema()",                          "Private",        "Creates all DB tables and indexes"),
    (34, "DatabaseManager",     "_seed_default_admin()",                         "Private",        "Inserts default admin account if none exists"),
    (35, "DatabaseManager",     "_seed_demo_items_if_empty()",                   "Private",        "Seeds 10 demo items on first run"),
    (36, "DatabaseManager",     "_log(connection, username, action, details)",   "Static method",  "Inserts a row into activity_logs"),
    (37, "DatabaseManager",     "authenticate_admin(username, password)",        "Instance",       "Returns Admin object or None"),
    (38, "DatabaseManager",     "log_activity(username, action, details)",       "Instance",       "Public activity log writer"),
    (39, "DatabaseManager",     "list_public_items(...)",                        "Instance",       "Filtered public catalog (safe columns only)"),
    (40, "DatabaseManager",     "get_public_item(item_id)",                      "Instance",       "Single public item"),
    (41, "DatabaseManager",     "get_public_locations()",                        "Instance",       "Distinct locations for filter UI"),
    (42, "DatabaseManager",     "list_admin_items(...)",                         "Instance",       "Full admin item list joined with finder info"),
    (43, "DatabaseManager",     "get_admin_item(item_id)",                       "Instance",       "Full admin item detail"),
    (44, "DatabaseManager",     "find_possible_duplicate_item(...)",             "Instance",       "Checks for duplicate item before adding"),
    (45, "DatabaseManager",     "create_found_item(finder_data, item_data, ...)" ,"Instance",      "Atomically inserts finder + item records"),
    (46, "DatabaseManager",     "update_item(item_id, item_data, admin)",        "Instance",       "Updates editable item fields"),
    (47, "DatabaseManager",     "archive_item(item_id, admin)",                  "Instance",       "Sets item status to 'Archived'"),
    (48, "DatabaseManager",     "create_claim(claim_data)",                      "Instance",       "Inserts claim, updates item to 'Pending Claim'"),
    (49, "DatabaseManager",     "list_claims(status, search)",                   "Instance",       "Filtered admin claims list"),
    (50, "DatabaseManager",     "get_claim(claim_row_id)",                       "Instance",       "Single claim detail"),
    (51, "DatabaseManager",     "update_claim_status(...)",                      "Instance",       "Approve / reject / request verification"),
    (52, "DatabaseManager",     "mark_returned(...)",                            "Instance",       "Finalizes return, sets item to 'Returned'"),
    (53, "DatabaseManager",     "list_finders(search)",                          "Instance",       "All finder records"),
    (54, "DatabaseManager",     "get_finder(finder_id)",                         "Instance",       "Single finder with their submitted items"),
    (55, "DatabaseManager",     "list_returns(search)",                          "Instance",       "All returned item records"),
    (56, "DatabaseManager",     "list_activity(search)",                         "Instance",       "Activity log entries (max 1000)"),
    (57, "DatabaseManager",     "dashboard_stats()",                             "Instance",       "Returns count dict for dashboard"),
    (58, "DatabaseManager",     "report_summary()",                              "Instance",       "Returns category/status breakdown dicts"),
    # AdminController
    (59, "AdminController",     "__init__(database)",                            "Instance",       "Stores DatabaseManager reference"),
    (60, "AdminController",     "login(username, password)",                     "Instance",       "Authenticates and logs activity"),
    (61, "AdminController",     "logout(admin)",                                 "Instance",       "Logs session end"),
    # ItemController
    (62, "ItemController",      "__init__(database)",                            "Instance",       "Stores DatabaseManager reference"),
    (63, "ItemController",      "public_catalog(...)",                           "Instance",       "Returns filtered public item list"),
    (64, "ItemController",      "public_item(item_id)",                          "Instance",       "Returns single public item"),
    (65, "ItemController",      "public_locations()",                            "Instance",       "Returns location filter options"),
    (66, "ItemController",      "admin_items(...)",                              "Instance",       "Returns full admin item list"),
    (67, "ItemController",      "admin_item(item_id)",                           "Instance",       "Returns admin item detail"),
    (68, "ItemController",      "possible_duplicate(item)",                      "Instance",       "Checks for duplicate before adding"),
    (69, "ItemController",      "add_found_item(finder, item, admin_username)",  "Instance",       "Validates and saves a new found item"),
    (70, "ItemController",      "update_item(item_id, values, admin_username)",  "Instance",       "Validates and updates item fields"),
    (71, "ItemController",      "archive_item(item_id, admin_username)",         "Instance",       "Archives an item"),
    (72, "ItemController",      "stats()",                                       "Instance",       "Returns dashboard count dict"),
    (73, "ItemController",      "report_summary()",                              "Instance",       "Returns report breakdown data"),
    # ClaimController
    (74, "ClaimController",     "__init__(database)",                            "Instance",       "Stores DatabaseManager reference"),
    (75, "ClaimController",     "submit_claim(claim)",                           "Instance",       "Validates and submits a claim"),
    (76, "ClaimController",     "list_claims(status, search)",                   "Instance",       "Returns filtered claims list"),
    (77, "ClaimController",     "claim(claim_row_id)",                           "Instance",       "Returns single claim detail"),
    (78, "ClaimController",     "review(claim_row_id, status, admin, notes)",    "Instance",       "Approve / reject / request verification"),
    (79, "ClaimController",     "mark_returned(claim_row_id, values, admin)",    "Instance",       "Finalizes item return"),
    (80, "ClaimController",     "stats()",                                       "Instance",       "Returns dashboard count dict"),
    # LostFoundSystem
    (81, "LostFoundSystem",     "__init__(database_path)",                       "Instance",       "Creates DatabaseManager + all three controllers"),
    # LostFoundApp
    (82, "LostFoundApp",        "__init__()",                                    "Instance",       "Sets up window, theme, system, and shows home"),
    (83, "LostFoundApp",        "_show(view_type, *args)",                       "Private",        "Destroys old view, renders new one"),
    (84, "LostFoundApp",        "show_public_home()",                            "Instance",       "Navigates to public catalog"),
    (85, "LostFoundApp",        "show_admin_login()",                            "Instance",       "Navigates to admin login screen"),
    (86, "LostFoundApp",        "show_admin_dashboard(admin)",                   "Instance",       "Navigates to admin dashboard"),
    (87, "LostFoundApp",        "logout_admin()",                                "Instance",       "Logs out admin and returns to public home"),
    (88, "LostFoundApp",        "_on_close()",                                   "Private",        "Confirms before closing the application"),
    # AdminLoginView
    (89, "AdminLoginView",      "__init__(master, app)",                         "Instance",       "Builds login UI with form fields"),
    (90, "AdminLoginView",      "_login()",                                      "Private",        "Reads fields, calls login, navigates on success"),
    # ScrollableFrame
    (91, "ScrollableFrame",     "__init__(master, background)",                  "Instance",       "Builds canvas + scrollbar layout"),
    (92, "ScrollableFrame",     "_update_scroll_region()",                       "Private",        "Updates canvas scroll bounds"),
    (93, "ScrollableFrame",     "_resize_inner(event)",                          "Private",        "Resizes inner frame to match canvas width"),
    (94, "ScrollableFrame",     "_enable_wheel()",                               "Private",        "Binds mousewheel scroll events"),
    (95, "ScrollableFrame",     "_disable_wheel()",                              "Private",        "Unbinds mousewheel scroll events"),
    (96, "ScrollableFrame",     "_on_mousewheel(event)",                         "Private",        "Windows scroll handler"),
    (97, "ScrollableFrame",     "_on_linux_wheel(event)",                        "Private",        "Linux scroll handler"),
    (98, "ScrollableFrame",     "_on_destroy(event)",                            "Private",        "Cleans up bindings on widget destroy"),
    (99, "ScrollableFrame",     "scroll_to(widget)",                             "Instance",       "Scrolls canvas to a specific child widget"),
]

# Group color by class
class_colors = {
    "PublicUser": "EAF1F8", "Admin": "EAF1F8",
    "SchoolPerson": "E9F5ED", "Student": "E9F5ED", "Teacher": "E9F5ED", "UtilityStaff": "E9F5ED",
    "Item": "FFF3D9", "Finder": "FFF3D9", "Claim": "FFF3D9",
    "DatabaseManager": "FCEBEC",
    "AdminController": "F4F7FA", "ItemController": "F4F7FA",
    "ClaimController": "F4F7FA", "LostFoundSystem": "F4F7FA",
    "LostFoundApp": "E4F4F2", "AdminLoginView": "E4F4F2", "ScrollableFrame": "E4F4F2",
}

for i, row in enumerate(methods_data, 3):
    for col, val in enumerate(row, 1):
        ws_methods.cell(row=i, column=col, value=val)
    cls = row[1]
    bg = class_colors.get(cls, WHITE)
    ws_methods.row_dimensions[i].height = 18
    for cell in ws_methods[i]:
        cell.fill      = fill(bg)
        cell.font      = cell_font()
        cell.alignment = wrap_center("left", "center", wrap=True)
        cell.border    = border()

set_col_widths(ws_methods, [5, 20, 44, 16, 52])


# ══════════════════════════════════════════════════════════════════════════════
# SHEET 4 — OOP CONCEPTS SUMMARY
# ══════════════════════════════════════════════════════════════════════════════
ws_oop = wb.create_sheet("OOP Concepts")

title_row(ws_oop, "Lost & Found System — OOP Concepts Applied", 3, bg=NAVY)

headers = ["OOP Concept", "Where Applied", "Example"]
for col, h in enumerate(headers, 1):
    ws_oop.cell(row=2, column=col, value=h)
style_header_row(ws_oop, 2, TEAL)

oop_data = [
    ("Inheritance",         "Admin → User, Student/Teacher/UtilityStaff → SchoolPerson, All Views → ttk.Frame",
                            "class Admin(User): ..."),
    ("Abstraction",         "SchoolPerson is an ABC with abstract properties role and affiliation",
                            "class SchoolPerson(ABC): @abstractmethod def role(self)"),
    ("Encapsulation",       "Protected _fields in SchoolPerson; Item.public_record() hides private data",
                            "self._full_name; Item.public_record() omits storage_location and private_notes"),
    ("Polymorphism",        "Finder.person and Claim.claimant accept any SchoolPerson subtype; to_record() dispatches correctly",
                            "finder.person.to_record() works for Student, Teacher, or UtilityStaff"),
    ("Composition",         "Finder and Claim contain a SchoolPerson object; LostFoundSystem composes all three controllers",
                            "class Finder: person: SchoolPerson"),
    ("Factory Function",    "make_school_person(role, data) returns the correct subtype",
                            "make_school_person('Student', data) → Student(...)"),
    ("Dataclasses",         "Item, Finder, Claim, User use @dataclass for auto-generated __init__, __repr__",
                            "@dataclass class Item: name: str; category: str; ..."),
    ("Class Methods",       "Admin.hash_password, Admin.verify_password, Item.from_mapping",
                            "@classmethod def hash_password(cls, password, salt)"),
    ("Static Methods",      "DatabaseManager._now(), DatabaseManager._row_dict(), DatabaseManager._log()",
                            "@staticmethod def _now() -> str"),
    ("Custom Exceptions",   "DuplicateClaimError and ItemUnavailableError extend ValueError",
                            "class DuplicateClaimError(ValueError): ..."),
    ("Facade Pattern",      "LostFoundSystem wraps DatabaseManager + 3 controllers as one entry point",
                            "app.system.items.add_found_item(...)"),
    ("Context Manager",     "DatabaseManager._connection() manages SQLite connection lifecycle",
                            "@contextmanager def _connection(self)"),
]

for i, row in enumerate(oop_data, 3):
    bg = LIGHT_GRAY if i % 2 == 0 else WHITE
    for col, val in enumerate(row, 1):
        c = ws_oop.cell(row=i, column=col, value=val)
        c.fill      = fill(bg)
        c.font      = cell_font(bold=(col == 1))
        c.alignment = wrap_center("left", "center", wrap=True)
        c.border    = border()
    ws_oop.row_dimensions[i].height = 28

set_col_widths(ws_oop, [22, 52, 52])

# ─── Save ─────────────────────────────────────────────────────────────────────
output_path = r"c:\Modules\EDOOP\lost-and-found-system\OOP_Summary_LostFoundSystem.xlsx"
wb.save(output_path)
print(f"Excel file saved to: {output_path}")
